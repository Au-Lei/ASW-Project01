"""Extract text from uploaded documents without retaining originals."""

import io
import locale
import os
import subprocess
import tempfile
from pathlib import Path

import openpyxl
import pdfplumber
from docx import Document

from app.state import RUNTIME_TMP_DIR


def decode_office_output(value: bytes | str) -> str:
    """Windows PowerShell may emit UTF-8 or the active ANSI code page."""
    if isinstance(value, str):
        return value
    try:
        return value.decode("utf-8")
    except UnicodeDecodeError:
        return value.decode(locale.getpreferredencoding(False), errors="replace")


def read_document(name: str, data: bytes) -> str:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    if suffix == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        parts.extend(" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
        return "\n".join(parts)
    if suffix == ".xlsx":
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            return "\n".join(" | ".join(str(cell.value) for cell in row if cell.value is not None) for sheet in wb for row in sheet)
        finally:
            wb.close()
    if suffix in {".doc", ".rtf", ".xls"}:
        if os.name != "nt":
            from app.tools.linux_office import to_pdf
            with pdfplumber.open(io.BytesIO(to_pdf(name, data))) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        temporary_root = RUNTIME_TMP_DIR
        temporary_root.mkdir(exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=temporary_root, suffix=suffix, delete=False) as temporary_file:
            temporary_file.write(data)
            source = Path(temporary_file.name)
        try:
            result = subprocess.run(
                ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(Path(__file__).with_name("legacy_extract.ps1")), "-Path", str(source)],
                capture_output=True, timeout=45, check=False,
            )
            output = decode_office_output(result.stdout)
            error = decode_office_output(result.stderr)
            if not output.strip():
                raise ValueError("旧版 Office 文件读取失败：" + (error.strip()[:250] or "未提取到文字"))
            return output
        finally:
            source.unlink(missing_ok=True)
    raise ValueError("支持 PDF、DOCX、XLSX，以及安装 Office 后的 DOC、RTF、XLS")
