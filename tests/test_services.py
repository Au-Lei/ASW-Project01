"""Contract and full local HTTP workflow tests using synthetic documents."""

import base64
import io
import json
import tempfile
import time
import unittest
from pathlib import Path
from threading import Thread
from urllib.error import HTTPError
from urllib.request import ProxyHandler, Request, build_opener
from unittest.mock import patch

import openpyxl
from docx import Document as WordDocument

from app import diagnostics
from app.business import review_config
from app.graph import extract
from app.models import FieldResult
from app.nodes.normalize import normalize_fields
from app.services.export import export_review
from app.services.requests import decode_extraction_request
from app.state import FIELDS, PROJECT_ROOT
from app.tools.document_reader import read_document
from app.web import Handler, ThreadingHTTPServer
from app import security


class ServiceTests(unittest.TestCase):
    def test_normalization_retains_raw_value_and_is_idempotent(self):
        fields = {key: FieldResult().to_dict() for key in FIELDS}
        fields["sailing_date"]["value"] = "08-Sep-2026"
        fields["sailing_date"].pop("raw_value")
        fields["destination"]["value"] = "HAMAD;QATAR"
        fields["destination"].pop("raw_value")
        normalized = normalize_fields(fields)
        self.assertEqual(normalized["sailing_date"]["value"], "9.8")
        self.assertEqual(normalized["sailing_date"]["raw_value"], "08-Sep-2026")
        self.assertEqual(normalized["destination"]["value"], "HAMAD")
        self.assertEqual(normalize_fields(normalized), normalized)
        self.assertEqual(fields["destination"]["value"], "HAMAD;QATAR")

    def test_pipeline_does_not_mutate_uploaded_documents(self):
        docs = [{"name": "notice.pdf", "role": "入货通知", "data": b"demo"}]
        with patch("app.graph.read_document", return_value="船期: 2026/9/8\n目的港: HAMAD,QATAR"):
            result = extract(docs, "rules")
        self.assertNotIn("text", docs[0])
        self.assertEqual(result["fields"]["destination"]["value"], "HAMAD")
        with self.assertRaises(ValueError):
            extract(docs, "unknown")

    def test_request_validation_rejects_malformed_input(self):
        for document in [None, {"role": "入货通知", "name": "x.exe", "data": "YQ=="},
                         {"role": "入货通知", "name": "x.pdf", "data": "!"},
                         {"role": "入货通知", "name": "x.pdf", "data": 123}]:
            with self.subTest(document=document), self.assertRaises(ValueError):
                decode_extraction_request({"documents": [document]})

    def test_export_validates_review_before_writing(self):
        with self.assertRaises(ValueError):
            export_review({"values": {"vessel": []}})
        with self.assertRaises(ValueError):
            export_review({"confirmed_aliases": "not-a-list"})

    def test_review_configuration_has_all_fields_and_staff(self):
        config = review_config()
        self.assertEqual(tuple(config["fields"]), FIELDS)
        self.assertEqual(set(config["manual_fields"]), {"carrier", "sales", "customer_service", "contract"})
        self.assertIn("ROBIN", config["staff_names"]["sales"])
        config["staff_names"]["sales"].clear()
        self.assertTrue(review_config()["staff_names"]["sales"])

    def test_docx_reader_uses_real_paragraphs_and_tables(self):
        document = WordDocument()
        document.add_paragraph("Vessel: DEMO SHIP")
        row = document.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = "Port of Discharge"
        row.cells[1].text = "HAMAD"
        stream = io.BytesIO()
        document.save(stream)
        text = read_document("synthetic.docx", stream.getvalue())
        self.assertIn("DEMO SHIP", text)
        self.assertIn("Port of Discharge | HAMAD", text)

    def test_diagnostics_omit_exception_values(self):
        logger = diagnostics.logging.getLogger("asw.diagnostics")
        previous = logger.handlers[:]
        previous_level, previous_propagate = logger.level, logger.propagate
        logger.handlers = []
        try:
            temporary_root = PROJECT_ROOT / ".runtime_tmp"
            temporary_root.mkdir(exist_ok=True)
            with tempfile.TemporaryDirectory(dir=temporary_root) as directory:
                with patch.object(diagnostics, "LOG_DIR", Path(directory) / "logs"):
                    try:
                        raise ValueError("SECRET_KEY_AND_CUSTOMER_CONTENT")
                    except ValueError as error:
                        diagnostics.record_error("test", error, job_id="test-job")
                    content = (Path(directory) / "logs" / "app.log").read_text(encoding="utf-8")
                # Traceback source lines can contain literals; diagnostic output must not.
                self.assertNotIn("SECRET_KEY_AND_CUSTOMER_CONTENT", content)
                self.assertIn("test-job", content)
                for handler in logger.handlers:
                    handler.close()
                logger.handlers = []
        finally:
            for handler in logger.handlers:
                handler.close()
            logger.handlers = previous
            logger.setLevel(previous_level)
            logger.propagate = previous_propagate


class HttpWorkflowTests(unittest.TestCase):
    def setUp(self):
        scratch = PROJECT_ROOT / ".runtime_tmp"
        scratch.mkdir(exist_ok=True)
        self.temp_dir = tempfile.TemporaryDirectory(dir=scratch)
        self.db_patch = patch.object(security, "DB_PATH", Path(self.temp_dir.name) / "security.sqlite3")
        self.key_patch = patch.object(security, "KEY_PATH", Path(self.temp_dir.name) / "security.key")
        self.db_patch.start(); self.key_patch.start()
        security.create_user("tester", "long-test-password-123")
        class QuietHandler(Handler):
            def log_message(self, *args):
                pass
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), QuietHandler)
        self.thread = Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.base = f"http://127.0.0.1:{self.server.server_port}"
        self.opener = build_opener(ProxyHandler({}))
        with self.request("/api/login", {"username": "tester", "password": "long-test-password-123"}) as response:
            login_data = json.load(response)
            self.cookie = response.headers["Set-Cookie"].split(";", 1)[0]
            self.csrf = login_data["csrf"]

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=2)
        self.db_patch.stop(); self.key_patch.stop(); self.temp_dir.cleanup()

    def request(self, path, payload=None):
        data = json.dumps(payload).encode() if payload is not None else None
        headers = {"Content-Type": "application/json"}
        if hasattr(self, "cookie"):
            headers.update({"Cookie": self.cookie, "X-CSRF-Token": self.csrf})
        return self.opener.open(Request(self.base + path, data=data, headers=headers), timeout=5)

    def test_static_assets_and_request_boundaries(self):
        with self.request("/") as response:
            html = response.read().decode()
        for name in ("api.js", "ui.js", "review.js", "memory.js", "app.js", "styles.css"):
            self.assertIn("/static/" + name, html)
            with self.request("/static/" + name) as response:
                self.assertTrue(response.read())
        with self.assertRaises(HTTPError) as error:
            self.request("/static/../state.py")
        self.assertEqual(error.exception.code, 404)
        error.exception.close()

    def test_http_auth_csrf_and_cross_account_job_denial(self):
        with self.assertRaises(HTTPError) as error:
            self.opener.open(self.base + "/api/stats", timeout=5)
        self.assertEqual(error.exception.code, 401)
        error.exception.close()
        with self.assertRaises(HTTPError) as error:
            self.opener.open(Request(self.base + "/api/export", data=b"{}",
                                     headers={"Content-Type": "application/json", "Cookie": self.cookie}), timeout=5)
        self.assertEqual(error.exception.code, 403)
        error.exception.close()
        from app.security import CURRENT_USER, create_user
        from app.jobs import start_job
        create_user("other", "other-account-password")
        context = CURRENT_USER.set(1)
        try:
            with patch("app.graph.extract", return_value={"fields": {field: {"value": ""} for field in FIELDS}}):
                job_id = start_job([{"name": "fake.pdf", "role": "入货通知", "data": b"demo"}], "rules")
        finally:
            CURRENT_USER.reset(context)
        login_body = json.dumps({"username": "other", "password": "other-account-password"}).encode()
        with self.opener.open(Request(self.base + "/api/login", data=login_body,
                                      headers={"Content-Type": "application/json"}), timeout=5) as response:
            other_cookie = response.headers["Set-Cookie"].split(";", 1)[0]
        with self.assertRaises(HTTPError) as error:
            self.opener.open(Request(self.base + "/api/extract/status?job_id=" + job_id,
                                     headers={"Cookie": other_cookie}), timeout=5)
        self.assertEqual(error.exception.code, 404)
        error.exception.close()
        with self.assertRaises(HTTPError) as error:
            self.request("/api/extract/start", [])
        self.assertEqual(json.load(error.exception)["error_code"], "invalid_request")
        error.exception.close()

    def test_xlsx_upload_progress_review_and_export(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for line in ["提单号: DEMO12345", "船名/航次: DEMO SHIP/001E", "船期: 2026/9/8",
                     "箱量: 1X20GP", "目的港: HAMAD,QATAR", "装货港: QINGDAO"]:
            sheet.append([line])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        payload = {"mode": "rules", "documents": [{"name": "synthetic.xlsx", "role": "入货通知",
                   "data": base64.b64encode(stream.getvalue()).decode()}]}
        with self.request("/api/extract/start", payload) as response:
            self.assertEqual(response.status, 202)
            job_id = json.load(response)["job_id"]
        deadline = time.monotonic() + 5
        while True:
            with self.request("/api/extract/status?job_id=" + job_id) as response:
                job = json.load(response)
            if job["status"] != "running" or time.monotonic() > deadline:
                break
            time.sleep(0.01)
        self.assertEqual(job["status"], "completed", job)
        fields = job["result"]["fields"]
        self.assertEqual(fields["origin"]["value"], "QD")
        self.assertEqual(fields["destination"]["value"], "HAMAD")
        values = {key: entry["value"] for key, entry in fields.items()}
        values.update(sales="ROBIN", customer_service="SISSIE", carrier="DEMO", contract="TEST")
        with self.request("/api/export", {"values": values}) as response:
            exported = openpyxl.load_workbook(io.BytesIO(response.read()))
        try:
            self.assertEqual(exported.active["B8"].value, "DEMO12345")
            self.assertEqual(exported.active["B10"].value, "9.8")
            self.assertEqual(exported.active["B39"].value, "ROBIN")
        finally:
            exported.close()
